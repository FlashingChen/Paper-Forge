# -*- coding: utf-8 -*-
"""Regression tests for the docx -> preview JSON extractor.

These tests are the enforcement of the preview's fidelity contract:

  * the text the preview shows is byte-for-byte the text python-docx reads out
    of the same file (so "预览和下载内容一致" is tested, not promised);
  * the JSON key names match src/lib/preview-types.ts (a rename fails here);
  * the counters in `stats` are decidable facts, recomputed from the payload.
"""
import contextlib
import io
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "agent" / "snippets"))
sys.path.insert(0, str(ROOT / "scripts"))
import docx_helpers as h  # noqa: E402
import docx_preview as p  # noqa: E402
from docx import Document  # noqa: E402

SAMPLE = ROOT / "reference" / "sample-worksheet.docx"
ANSWER = "\uFF3F"

PX_PER_MM = 96.0 / 25.4


def build_doc(path: Path, merge: bool = True):
    """A spec-compliant sheet: title, body, answer lines, a table."""
    doc = h.new_doc()
    h.add_centered(doc, "春日示例", size_pt=12, bold=True)
    h.add_para(doc, "1.请结合材料阐述理由。（3分）", indent_chars=2)
    h.add_answer_lines(doc, 2, prefix="答：")
    h.add_para(doc, "2.解释下列加点的词。（2分）")
    table = h.add_table(
        doc,
        [["", "释义"], ["趋", "（1）" + h.answer_fill()], ["顾", "（2）" + h.answer_fill(5)]],
        widths_mm=[80, 80],
    )
    if merge:
        # Exercises w:vMerge (restart + continuation) in the extractor. Note
        # python-docx MOVES the lower cell's content into the merged top cell,
        # which is why the verbatim text comparison uses a merge-free document.
        table.cell(1, 0).merge(table.cell(2, 0))
    doc.save(str(path))
    return doc


def iter_blocks(blocks):
    for block in blocks:
        yield block
        if block["type"] == "table":
            for row in block["rows"]:
                for cell in row["cells"]:
                    yield from iter_blocks(cell["blocks"])


def block_text(block) -> str:
    return "".join(run["text"] for run in block["runs"])


def cell_texts(document) -> list:
    """Rendered cell texts, paragraph-joined the way python-docx cell.text is."""
    texts = []
    for block in document["blocks"]:
        if block["type"] != "table":
            continue
        for row in block["rows"]:
            for cell in row["cells"]:
                if cell["skip"]:
                    continue
                texts.append("\n".join(block_text(b) for b in cell["blocks"]))
    return texts


class PreviewTest(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def built(self):
        """Path of a freshly built document (with a vertically merged cell)."""
        path = self.tmp / "result.docx"
        build_doc(path)
        return path

    def flat(self):
        """A merge-free document, for verbatim text comparison."""
        path = self.tmp / "flat.docx"
        build_doc(path, merge=False)
        return path

    # ------------------------------------------------------------- geometry --

    def test_page_geometry_is_a4_with_20_25mm_margins(self):
        for path in (SAMPLE, self.built()):
            with self.subTest(path=path.name):
                page = p.build_preview(str(path))["page"]
                self.assertAlmostEqual(page["widthPx"], 210 * PX_PER_MM, delta=1)
                self.assertAlmostEqual(page["heightPx"], 297 * PX_PER_MM, delta=1)
                self.assertAlmostEqual(page["paddingLeftPx"], 25 * PX_PER_MM, delta=1)
                self.assertAlmostEqual(page["paddingRightPx"], 25 * PX_PER_MM, delta=1)
                self.assertAlmostEqual(page["paddingTopPx"], 20 * PX_PER_MM, delta=1)
                self.assertAlmostEqual(page["paddingBottomPx"], 20 * PX_PER_MM, delta=1)

    # ------------------------------------------------------------ text parity --

    def test_text_matches_python_docx_verbatim(self):
        path = self.flat()
        preview = p.build_preview(str(path))
        doc = Document(str(path))

        shown = [
            block_text(b) for b in preview["blocks"] if b["type"] == "p"
        ]
        self.assertEqual(shown, [para.text for para in doc.paragraphs])

        cells = [
            cell.text for table in doc.tables for row in table.rows for cell in row.cells
        ]
        self.assertEqual(cell_texts(preview), cells)

    def test_sample_worksheet_round_trips(self):
        preview = p.build_preview(str(SAMPLE))
        doc = Document(str(SAMPLE))
        shown = [block_text(b) for b in preview["blocks"] if b["type"] == "p"]
        self.assertEqual(shown, [para.text for para in doc.paragraphs])
        self.assertIn(ANSWER, "".join(shown) + "".join(cell_texts(preview)))

    # --------------------------------------------------------------- styling --

    def test_answer_lines_are_kaiti_and_body_is_songti(self):
        preview = p.build_preview(str(self.built()))
        fonts = preview["stats"]["fonts"]
        self.assertEqual(fonts, ["宋体", "楷体"])

        answer_run = None
        body_run = None
        for block in iter_blocks(preview["blocks"]):
            if block["type"] != "p":
                continue
            for run in block["runs"]:
                if ANSWER in run["text"] and answer_run is None:
                    answer_run = run
                if "请结合材料" in run["text"] and body_run is None:
                    body_run = run
        self.assertIsNotNone(answer_run)
        self.assertIsNotNone(body_run)
        self.assertEqual(answer_run["font"], "楷体")
        self.assertEqual(body_run["font"], "宋体")
        self.assertEqual(body_run["sizePt"], 10.5)

    def test_first_line_indent_is_two_characters(self):
        preview = p.build_preview(str(self.built()))
        target = None
        for block in preview["blocks"]:
            if block["type"] == "p" and block_text(block).startswith("1.请结合材料"):
                target = block
        self.assertIsNotNone(target)
        # 2 characters at 10.5pt = 2 * 10.5 * 4/3 px.
        self.assertAlmostEqual(target["indentFirstLinePx"], 2 * 10.5 * 4 / 3, delta=1)
        self.assertEqual(target["align"], "justify")

    def test_character_based_first_line_indent(self):
        """Word writes w:firstLineChars for Chinese indents; honour it too."""
        from docx.oxml.ns import qn as dqn

        path = self.tmp / "chars.docx"
        doc = h.new_doc()
        paragraph = h.add_para(doc, "字符缩进的段落。", size_pt=12)
        ind = paragraph._p.get_or_add_pPr().get_or_add_ind()
        ind.set(dqn("w:firstLineChars"), "200")
        doc.save(str(path))

        preview = p.build_preview(str(path))
        block = next(b for b in preview["blocks"] if b["type"] == "p")
        # 200 hundredths of a character = 2 characters at 12pt.
        self.assertAlmostEqual(block["indentFirstLinePx"], 2 * 12 * 4 / 3, delta=1)

    # ---------------------------------------------------------------- tables --

    def test_table_spans_and_vertical_merge(self):
        preview = p.build_preview(str(self.built()))
        tables = [b for b in preview["blocks"] if b["type"] == "table"]
        self.assertEqual(len(tables), 1)
        table = tables[0]

        self.assertEqual(len(table["widthsPx"]), 2)
        for row in table["rows"]:
            self.assertEqual(sum(c["colSpan"] for c in row["cells"]), 2)

        merged_start = table["rows"][1]["cells"][0]
        merged_continue = table["rows"][2]["cells"][0]
        self.assertEqual(merged_start["rowSpan"], 2)
        self.assertFalse(merged_start["skip"])
        self.assertTrue(merged_continue["skip"])
        self.assertEqual(block_text(merged_start["blocks"][0]), "趋")
        self.assertEqual(merged_continue["blocks"], [])

    # ----------------------------------------------------------------- stats --

    def test_stats_are_decidable_from_the_payload(self):
        preview = p.build_preview(str(self.built()))
        stats = preview["stats"]

        paragraphs = [b for b in preview["blocks"] if b["type"] == "p"]
        self.assertEqual(stats["paragraphs"], len(paragraphs))
        self.assertEqual(stats["tables"], len([b for b in preview["blocks"] if b["type"] == "table"]))

        all_paragraphs = [b for b in iter_blocks(preview["blocks"]) if b["type"] == "p"]
        text = "".join(block_text(b) for b in all_paragraphs)
        self.assertEqual(stats["answerChars"], text.count(ANSWER))
        self.assertEqual(
            stats["answerLineRows"],
            len([b for b in all_paragraphs if ANSWER in block_text(b)]),
        )
        self.assertEqual(stats["textChars"], len(text) - text.count(ANSWER))

        cells = []
        for block in preview["blocks"]:
            if block["type"] == "table":
                for row in block["rows"]:
                    cells.extend([c for c in row["cells"] if not c["skip"]])
        self.assertEqual(stats["tableCells"], len(cells))

    def test_truncation_is_flagged(self):
        original = p.MAX_BLOCKS
        p.MAX_BLOCKS = 2
        try:
            preview = p.build_preview(str(self.built()))
        finally:
            p.MAX_BLOCKS = original
        self.assertTrue(preview["truncated"])
        self.assertTrue(any("截断" in w for w in preview["warnings"]))
        self.assertEqual(len(preview["blocks"]), 2)

    def test_compliant_document_produces_no_warnings(self):
        preview = p.build_preview(str(self.built()))
        self.assertEqual(preview["warnings"], [])
        self.assertFalse(preview["truncated"])

    def test_missing_fonts_are_reported(self):
        path = self.tmp / "bare.docx"
        doc = Document()
        doc.add_paragraph("没有任何显式字体设置的段落")
        doc.save(str(path))
        preview = p.build_preview(str(path))
        self.assertTrue(
            any("没有显式指定字体" in w for w in preview["warnings"]),
            preview["warnings"],
        )

    def test_content_control_wrapper_keeps_its_text(self):
        """Word wraps paragraphs in w:sdt all the time; do not drop them."""
        from docx.oxml import OxmlElement

        path = self.tmp / "sdt.docx"
        doc = h.new_doc()
        h.add_para(doc, "普通段落")
        inside = h.add_para(doc, "内容控件里的段落")
        body = inside._p.getparent()
        body.remove(inside._p)
        wrapper = OxmlElement("w:sdt")
        content = OxmlElement("w:sdtContent")
        content.append(inside._p)
        wrapper.append(content)
        body.insert(0, wrapper)
        doc.save(str(path))

        preview = p.build_preview(str(path))
        shown = [block_text(b) for b in preview["blocks"] if b["type"] == "p"]
        self.assertIn("内容控件里的段落", shown)
        self.assertEqual(preview["warnings"], [])

    def test_unknown_text_bearing_structure_is_reported(self):
        """Losing text silently would break the core preview promise."""
        from docx.oxml import OxmlElement

        path = self.tmp / "weird.docx"
        doc = h.new_doc()
        h.add_para(doc, "正常段落")
        body = doc.element.body
        weird = OxmlElement("w:notAThing")
        text = OxmlElement("w:t")
        text.text = "藏在认不出的结构里的字"
        weird.append(text)
        body.insert(0, weird)
        doc.save(str(path))

        preview = p.build_preview(str(path))
        self.assertTrue(
            any("认不出" in w for w in preview["warnings"]),
            preview["warnings"],
        )

    # -------------------------------------------------------------- contract --

    def test_json_key_names_match_the_client_contract(self):
        preview = p.build_preview(str(self.built()))
        self.assertEqual(
            set(preview),
            {"ok", "page", "stats", "blocks", "truncated", "warnings"},
        )
        self.assertEqual(
            set(preview["page"]),
            {
                "widthPx",
                "heightPx",
                "paddingTopPx",
                "paddingRightPx",
                "paddingBottomPx",
                "paddingLeftPx",
            },
        )
        self.assertEqual(
            set(preview["stats"]),
            {
                "paragraphs",
                "tables",
                "tableCells",
                "answerLineRows",
                "answerChars",
                "textChars",
                "fonts",
            },
        )

        paragraph = next(b for b in preview["blocks"] if b["type"] == "p")
        self.assertEqual(
            set(paragraph),
            {
                "type",
                "align",
                "indentFirstLinePx",
                "indentLeftPx",
                "spaceBeforePx",
                "spaceAfterPx",
                "lineRule",
                "line",
                "runs",
            },
        )
        self.assertEqual(
            set(paragraph["runs"][0]),
            {"text", "font", "sizePt", "bold", "italic", "color", "underline"},
        )

        table = next(b for b in preview["blocks"] if b["type"] == "table")
        self.assertEqual(set(table), {"type", "widthsPx", "rows"})
        self.assertEqual(set(table["rows"][0]), {"cells"})
        self.assertEqual(
            set(table["rows"][0]["cells"][0]),
            {"colSpan", "rowSpan", "skip", "blocks"},
        )

    def test_payload_is_json_serialisable_without_escapes(self):
        preview = p.build_preview(str(self.built()))
        encoded = json.dumps(preview, ensure_ascii=False)
        self.assertIn("宋体", encoded)
        self.assertEqual(json.loads(encoded), preview)

    # ----------------------------------------------------------- error paths --

    def test_main_reports_bad_input_as_json(self):
        not_a_docx = self.tmp / "photo.jpg"
        not_a_docx.write_bytes(b"\xff\xd8\xff\xe0not a zip")

        for argv in (
            ["docx_preview.py", str(not_a_docx)],
            ["docx_preview.py", str(self.tmp / "missing.docx")],
            ["docx_preview.py"],
            ["docx_preview.py", str(self.nameless_zip())],
        ):
            with self.subTest(argv=argv):
                buffer = io.StringIO()
                with contextlib.redirect_stdout(buffer):
                    code = p.main(argv)
                self.assertEqual(code, 2)
                payload = json.loads(buffer.getvalue())
                self.assertFalse(payload["ok"])
                self.assertTrue(payload["error"])

    def test_main_succeeds_on_a_real_document(self):
        buffer = io.StringIO()
        with contextlib.redirect_stdout(buffer):
            code = p.main(["docx_preview.py", str(self.built())])
        self.assertEqual(code, 0)
        self.assertTrue(json.loads(buffer.getvalue())["ok"])

    def nameless_zip(self) -> Path:
        """A valid zip that is not a Word document."""
        path = self.tmp / "empty.docx"
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr("hello.txt", "not a document")
        return path


if __name__ == "__main__":
    unittest.main()
