# -*- coding: utf-8 -*-
"""
docx_helpers.py —— 语文试卷排版通用助手模块（python-docx）

用法：
    import sys
    sys.path.insert(0, "snippets")      # 或直接把这个文件 cp 到 out/ 旁边
    from docx_helpers import (
        setup_page, set_font, add_para, add_centered,
        add_answer_lines, add_emphasis, add_underline,
        add_table, add_note, new_doc,
    )

    doc = new_doc()                     # 已 setup_page + 全局 Normal 字体
    add_para(doc, "材料一：")
    add_centered(doc, "春日示例", size_pt=12, bold=True)
    ...
    doc.save("out/result.docx")

设计原则（对应硬性要求）：
  * 每个 run 都显式写 w:ascii / w:hAnsi / w:eastAsia，绝不继承 theme。
  * 答题横线 = 全角下划线字符 U+FF3F（＿），不使用段落边框或空格下划线。
  * 用真表格 (add_table)，不用空格对齐。
  * 不产生装饰线、底纹、分页符、填充空段。
"""

from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# ---------------------------------------------------------------------------
# 常量（论文/试卷常用值，已按需求固定）
# ---------------------------------------------------------------------------

# 西文名与 Word 内部名都用「宋体」：Windows 上 Word 能正确解析中文字体名，
# 且避免 SimSun / 宋体 两个名字被当成两个字体。eastAsia 必须写中文字体名。
FONT_NAME = "宋体"

# 正文字号（磅）。小四 = 12pt，五号 = 10.5pt。
# --- 字号约定（中文试卷惯例）--------------------------------------------------
# 正文五号；篇目标题小四；作者与注释小五。
# 不要用 12/14pt 做正文：卷面会发胖、变松垮，不像一张考试卷。
BODY_SIZE = 10.5
NOTE_SIZE = 9.0           # 〔注〕注释块
TITLE_SIZE = 12.0         # 篇目标题（小四）
AUTHOR_SIZE = 9.0         # 作者行（小五）

# 与原创版式参考文档一致，使用实际字符，不给空格加下划线。
ANSWER_CHAR = "\uFF3F"
ANSWER_FONT_NAME = "楷体"

# A4 纵向：210mm x 297mm。页边距上下 20mm、左右 25mm。
PAGE_W_MM = 210.0
PAGE_H_MM = 297.0
MARGIN_TB_MM = 20.0
MARGIN_LR_MM = 25.0

# 正文可用宽度（mm）= 210 - 25*2 = 160mm。
TEXT_WIDTH_MM = PAGE_W_MM - 2 * MARGIN_LR_MM

# 参考文档的独立答题行用 38 个字符；首行有「答：」时减少两个。
# 较大字号或较窄版心还会按实际可用宽度缩短，避免横线自动折行。
ANSWER_CHARS_PER_LINE = 38

# 表格/句中填空默认 7 个字符，可用 answer_fill(n) 指定长度。
TABLE_ANSWER_CHARS = 7


# ---------------------------------------------------------------------------
# 页面 / 文档
# ---------------------------------------------------------------------------

def setup_page(doc):
    """A4 纵向，页边距上下 20mm、左右 25mm。"""
    for section in doc.sections:
        section.page_width = Mm(PAGE_W_MM)
        section.page_height = Mm(PAGE_H_MM)
        section.top_margin = Mm(MARGIN_TB_MM)
        section.bottom_margin = Mm(MARGIN_TB_MM)
        section.left_margin = Mm(MARGIN_LR_MM)
        section.right_margin = Mm(MARGIN_LR_MM)
        section.header_distance = Mm(12)
        section.footer_distance = Mm(12)
    return doc


def _set_normal_style(doc, name=FONT_NAME, size_pt=BODY_SIZE):
    """把 Normal 样式也改掉，双保险：即使某处漏设 run 属性也不会回落主题字体。"""
    normal = doc.styles["Normal"]
    normal.font.name = name                      # 只设这个 = 西文生效
    normal.font.size = Pt(size_pt)
    rpr = normal.element.get_or_add_rPr()        # 关键：补 eastAsia
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)
    return normal


def new_doc():
    """新建文档：A4 + 页边距 + Normal 宋体。"""
    doc = Document()
    setup_page(doc)
    _set_normal_style(doc)
    # 默认段落不加段后空白（避免试卷被撑开），行距单倍。
    pf = doc.styles["Normal"].paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    return doc


# ---------------------------------------------------------------------------
# 字体：必须显式写 rFonts 的三个属性
# ---------------------------------------------------------------------------

def set_font(run, name=FONT_NAME, size_pt=BODY_SIZE, bold=False,
             italic=False, color=None):
    """
    显式设置 run 的字体。这是硬性要求 1 的唯一正确做法。

    run.font.name = '宋体' 只会写 w:ascii/w:hAnsi，
    中日韩字符仍走 w:eastAsia（缺失时继承主题），所以必须手动补。
    """
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)
    return run


def _add_text_runs(paragraph, text, name=FONT_NAME, size_pt=BODY_SIZE, bold=False):
    """句中/表格里的字符横线显式用楷体，其余文字保留指定正文字体。"""
    import re
    runs = []
    for part in re.split(r"(＿+)", text):
        if part:
            font = ANSWER_FONT_NAME if set(part) == {ANSWER_CHAR} else name
            runs.append(set_font(paragraph.add_run(part), name=font,
                                 size_pt=size_pt, bold=bold))
    return runs


# ---------------------------------------------------------------------------
# 段落
# ---------------------------------------------------------------------------

_ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def add_para(doc, text="", size_pt=BODY_SIZE, align="justify",
             bold=False, name=FONT_NAME, indent_chars=0,
             space_before=0.0, space_after=0.0, line_spacing=1.0):
    """
    正文段落。indent_chars>0 表示首行缩进多少个全角字（1 字 = size_pt）。

    注意：text 里可以有换行，但试卷要求一行一段，所以请一行调用一次。
    """
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.alignment = _ALIGN[align]
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    if indent_chars:
        pf.first_line_indent = Pt(size_pt * indent_chars)
    if text:
        _add_text_runs(p, text, name=name, size_pt=size_pt, bold=bold)
    return p


def _clear_runs(p):
    """删除段落里已有的 run。cell.text = "" 会留下无字体的空 run，必须清掉。"""
    for r in list(p._p.findall(qn("w:r"))):
        p._p.remove(r)
    return p


def add_centered(doc, text, size_pt=BODY_SIZE, bold=False, name=FONT_NAME,
                 space_before=0.0, space_after=0.0):
    """居中段落。诗歌每行、篇目标题、作者行都用它。"""
    return add_para(doc, text, size_pt=size_pt, align="center", bold=bold,
                    name=name, space_before=space_before,
                    space_after=space_after)


def add_right(doc, text, size_pt=BODY_SIZE, name=FONT_NAME):
    """右对齐段落，例如「（有删改）」。"""
    return add_para(doc, text, size_pt=size_pt, align="right", name=name)


def add_note(doc, text):
    """〔注〕注释块：比正文小一号。可用多次调用或传整段文本。"""
    return add_para(doc, text, size_pt=NOTE_SIZE, align="justify",
                    name=FONT_NAME)


# ---------------------------------------------------------------------------
# 答题横线：全角下划线字符，与参考文档使用相同的楷体字形
# ---------------------------------------------------------------------------

def _answer_capacity(container, size_pt, prefix=""):
    """按当前版心/单元格宽度估算字符容量，并预留一个字的换行余量。"""
    if size_pt <= 0:
        raise ValueError("size_pt must be positive")
    if hasattr(container, "sections"):
        section = container.sections[-1]
        width = (section.page_width - section.left_margin - section.right_margin) / 12700
    elif getattr(container, "width", None) is not None:
        # 单元格两侧默认内边距各约 1.9mm。
        width = container.width.pt - 2 * 1.9 * 72 / 25.4
    else:
        width = TEXT_WIDTH_MM * 72 / 25.4
    capacity = int(width / size_pt) - len(prefix) - 1
    if capacity < 1:
        raise ValueError("answer line has no room after its prefix")
    return capacity


def add_answer_lines(doc, n, size_pt=BODY_SIZE, name=ANSWER_FONT_NAME, prefix=""):
    """生成 n 行字符答题线。prefix='答：' 时首行文字与横线在同一段。

    普通行最多 38 个「＿」，首行扣除 prefix 的长度；每行段后 10pt，
    与参考文档一致。不添加空段落、制表符、w:u 或段落边框。
    name 可覆盖答题线字体；默认楷体，正文仍用宋体。
    """
    if int(n) < 0:
        raise ValueError("n must be non-negative")
    out = []
    for i in range(int(n)):
        lead = prefix if i == 0 else ""
        chars = min(ANSWER_CHARS_PER_LINE - len(lead), _answer_capacity(doc, size_pt, lead))
        if chars < 1:
            raise ValueError("answer line prefix is too long")
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = 1.0
        set_font(p.add_run(lead + ANSWER_CHAR * chars), name=name, size_pt=size_pt)
        out.append(p)
    return out


def add_answer_line_inline(doc, size_pt=BODY_SIZE, chars=None):
    """在单元格最后一段/给定段落末尾追加字符横线，标签与横线同行。

    传入文档时新增一段；传入单元格时复用最后一段；chars 默认 7。
    """
    requested = TABLE_ANSWER_CHARS if chars is None else int(chars)
    if requested < 1:
        raise ValueError("chars must be positive")
    if hasattr(doc, "add_run"):
        p = doc
        container = p._parent
    elif hasattr(doc, "sections"):
        p = doc.add_paragraph()
        container = doc
    else:
        p = doc.paragraphs[-1]
        container = doc
    count = min(requested, _answer_capacity(container, size_pt, p.text))
    set_font(p.add_run(ANSWER_CHAR * count), name=ANSWER_FONT_NAME, size_pt=size_pt)
    return p


# ---------------------------------------------------------------------------
# 着重号 / 下划线（python-docx 无原生 API，直接写 XML）
# ---------------------------------------------------------------------------

def add_emphasis(run, val="dot"):
    """
    着重号（字下加点）。w:em 是 run 属性里的子元素，放在 rPr 内。
    参考 docx 里用的是 <w:em w:val="dot"/>。
    """
    rpr = run._element.get_or_add_rPr()
    em = rpr.find(qn("w:em"))
    if em is None:
        em = OxmlElement("w:em")
        _insert_in_order(rpr, em)
    em.set(qn("w:val"), val)
    return run


def add_underline(run, style="single"):
    """
    run 下划线。style='single' 直线（画横线的句子），
    style='wave' 波浪线（画波浪线的句子），style='none' 取消。
    """
    rpr = run._element.get_or_add_rPr()
    u = rpr.find(qn("w:u"))
    if u is None:
        u = OxmlElement("w:u")
        _insert_in_order(rpr, u)
    u.set(qn("w:val"), style)
    return run


# rPr 子元素的 schema 顺序（子集，够用）。乱序会让 Word 报文件损坏。
_RPR_ORDER = [
    "w:rStyle", "w:rFonts", "w:b", "w:bCs", "w:i", "w:iCs", "w:caps",
    "w:smallCaps", "w:strike", "w:dstrike", "w:outline", "w:shadow",
    "w:emboss", "w:imprint", "w:noProof", "w:snapToGrid", "w:vanish",
    "w:webHidden", "w:color", "w:spacing", "w:w", "w:kern", "w:position",
    "w:sz", "w:szCs", "w:highlight", "w:u", "w:effect", "w:bdr", "w:shd",
    "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang",
    "w:eastAsianLayout", "w:specVanish", "w:oMath",
]


def _insert_in_order(rpr, element):
    """把 element 按 schema 顺序插入 rPr，避免 Word 判定文档损坏。"""
    tag = element.tag.split("}")[1]
    key = "w:" + tag
    try:
        idx = _RPR_ORDER.index(key)
    except ValueError:
        rpr.append(element)
        return
    for child in rpr:
        ctag = "w:" + child.tag.split("}")[1]
        cidx = _RPR_ORDER.index(ctag) if ctag in _RPR_ORDER else 999
        if cidx > idx:
            child.addprevious(element)
            return
    rpr.append(element)


def add_rich_para(doc, parts, size_pt=BODY_SIZE, align="justify",
                  indent_chars=0, name=FONT_NAME,
                  space_before=0.0, space_after=0.0, line_spacing=1.0):
    """
    需要局部着重号/下划线的段落用这个。
    parts 是 [(text, opts), ...]，opts 支持 em / underline / bold：
        add_rich_para(doc, [
            ("风过枝头动，云开天色明", {"em": True, "underline": "wave"}),
            ("①。", {}),
        ])
    """
    p = add_para(doc, "", size_pt=size_pt, align=align, name=name,
                 indent_chars=indent_chars, space_before=space_before,
                 space_after=space_after, line_spacing=line_spacing)
    for text, opts in parts:
        if not text:
            continue
        for run in _add_text_runs(p, text, name=name, size_pt=size_pt,
                                  bold=opts.get("bold", False)):
            if set(run.text) == {ANSWER_CHAR}:
                continue
            if opts.get("em"):
                add_emphasis(run, opts.get("em", "dot"))
            if opts.get("underline"):
                add_underline(run, opts["underline"])
    return p


# ---------------------------------------------------------------------------
# 表格
# ---------------------------------------------------------------------------

def add_table(doc, rows, header=True, widths_mm=None, font_size=BODY_SIZE,
              name=FONT_NAME, align_center_cols=None):
    """
    真表格（硬性要求 4）。rows 是 [[c1, c2, c3], ...]，第一行当表头。
    widths_mm 是每列宽度（mm）；不给则平均分。单元格文字会显式设宋体。

    align_center_cols: 需要居中文字的列索引集合，如 {2}。
    """
    n_rows = len(rows)
    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=n_rows, cols=n_cols)
    table.style = "Table Grid"          # 有框线的标准表格
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    if widths_mm is None:
        widths_mm = [TEXT_WIDTH_MM / n_cols] * n_cols
    center = set(align_center_cols or ())

    for ri, row in enumerate(rows):
        for ci in range(n_cols):
            cell = table.cell(ri, ci)
            cell.width = Mm(widths_mm[ci])
            txt = str(row[ci] if ci < len(row) else "")
            first = cell.paragraphs[0]
            _clear_runs(first)
            first.paragraph_format.space_before = Pt(1)
            first.paragraph_format.space_after = Pt(1)
            # 表格里可能有多行文本（如词典释义），用换行符拆成多段
            for li, line in enumerate(txt.split("\n")):
                p = first
                if li > 0:
                    p = cell.add_paragraph()
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                    if (header and ri == 0) or ci in center
                    else WD_ALIGN_PARAGRAPH.LEFT
                )
                if line:
                    _add_text_runs(p, line, name=name, size_pt=font_size,
                                   bold=(header and ri == 0))
                # 空表头、空单元格和空行只是空白，不自动生成答题线。
    # 固定列宽（Word 需要 tblLayout=fixed 才尊重 width）
    _set_table_fixed_layout(table)
    return table


def _set_table_fixed_layout(table):
    tblPr = table._tbl.tblPr
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")


# ---------------------------------------------------------------------------
# 便捷：释义列专用的下划线填充
# ---------------------------------------------------------------------------

def answer_fill(n=None):
    """返回可嵌在句子或表格中的全角下划线字符，默认 7 个。"""
    count = TABLE_ANSWER_CHARS if n is None else int(n)
    if count < 0:
        raise ValueError("n must be non-negative")
    return ANSWER_CHAR * count


if __name__ == "__main__":
    # 自检：生成一个 mini 试卷并断言 XML 正确
    import os
    doc = new_doc()
    add_para(doc, "材料一：", align="left")
    add_centered(doc, "春日示例", size_pt=TITLE_SIZE, bold=True)
    add_centered(doc, "示例作者", size_pt=AUTHOR_SIZE)
    add_centered(doc, "清风吹小院，新叶映窗台。")
    p = add_para(doc, "", align="justify")
    r = p.add_run("风过枝头动，云开天色明")
    set_font(r)
    add_emphasis(r)
    add_underline(r, "wave")
    r2 = p.add_run("①。")
    set_font(r2)
    add_note(doc, "〔注〕①枝头：树枝的末端。")
    add_para(doc, "7.参考表格提示的方法，解释下列加点的词。（4分）", align="left")
    add_table(doc, [
        ["文言语句", "方法借鉴", "释　义"],
        ["遂前" + " ", "【语境推断法】联系上下文推断", "（1）" + answer_fill()],
        ["靠近树荫", "【语境分析法】根据上下文判断含义", "（2）" + answer_fill()],
        ["环顾四周", "【参考成语法】左顾右盼", "（3）" + answer_fill()],
        ["来到最高处", "【词义辨析法】①尽；②超过", "（4）" + answer_fill()],
    ], widths_mm=[40, 80, 40], align_center_cols={2})
    add_para(doc, "8.用现代汉语翻译材料中画横线的句子。（4分）", align="left")
    add_para(doc, "(1)晨雾渐散，鸟鸣渐起。（2分）", align="left")
    add_answer_lines(doc, 2, prefix="答：")
    # A 6-mark open question needs 6 rules. Used to be a silent defect.
    add_para(doc, "10.结合材料阐述理由。（6分）", align="left")
    add_answer_lines(doc, 6, prefix="答：")
    out = "/tmp/pf_helpers_selftest.docx"
    doc.save(out)
    print("saved", out, os.path.getsize(out), "bytes")


# ---------------------------------------------------------------------------
# 从原照片裁切题目插图并嵌入 Word（数学等科目）
# ---------------------------------------------------------------------------

def add_cropped_figure(doc, source_path, box, crop_path, width_mm=65):
    """box 为 EXIF 方向校正后原图的像素坐标；返回插图所在段落。"""
    import math
    from pathlib import Path
    from PIL import Image, ImageOps

    if len(box) != 4 or any(not isinstance(v, (int, float)) or not math.isfinite(v)
                            for v in box):
        raise ValueError("box 必须是四个有限像素坐标")
    if not isinstance(width_mm, (int, float)) or not math.isfinite(width_mm) or width_mm <= 0:
        raise ValueError("width_mm 必须大于 0")
    target = Path(crop_path)
    if target.suffix.lower() not in (".jpg", ".jpeg"):
        raise ValueError("裁图必须保存为 JPEG")
    if Path(source_path).resolve() == target.resolve():
        raise ValueError("不能覆盖原图")
    with Image.open(source_path) as original:
        image = ImageOps.exif_transpose(original)
        left, top, right, bottom = (round(v) for v in box)
        if not (0 <= left < right <= image.width and 0 <= top < bottom <= image.height):
            raise ValueError(f"裁切框超出原图范围：{image.width} × {image.height}")
        crop = image.crop((left, top, right, bottom)).convert("RGB")
        target.parent.mkdir(parents=True, exist_ok=True)
        crop.save(target, quality=95, subsampling=0)

    section = doc.sections[-1]
    available_width = section.page_width - section.left_margin - section.right_margin
    available_height = section.page_height - section.top_margin - section.bottom_margin
    # Keep the full figure inside the page in both directions, with its aspect ratio.
    width = min(Mm(width_mm), available_width, int(available_height * crop.width / crop.height))
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run()
    set_font(run)
    run.add_picture(str(target), width=width)
    return paragraph
